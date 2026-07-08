from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "qua_trinh_nghien_cuu_neural_operator_vbts.pptx"

W, H = 13.333, 7.5

NAVY = RGBColor(18, 33, 55)
INK = RGBColor(37, 45, 54)
MUTED = RGBColor(92, 108, 126)
PAPER = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(218, 226, 235)
BLUE = RGBColor(32, 99, 155)
CYAN = RGBColor(0, 151, 167)
GREEN = RGBColor(54, 132, 104)
ORANGE = RGBColor(198, 117, 46)
RED = RGBColor(171, 71, 66)
PURPLE = RGBColor(108, 92, 160)
PALE_BLUE = RGBColor(232, 241, 250)
PALE_GREEN = RGBColor(234, 246, 240)
PALE_ORANGE = RGBColor(252, 242, 230)
PALE_RED = RGBColor(252, 235, 235)


def load_json(rel: str) -> dict:
    with (ROOT / rel).open("r", encoding="utf-8") as f:
        return json.load(f)


def fnum(x: float, nd: int = 2) -> str:
    return f"{x:.{nd}f}"


def pct(x: float, nd: int = 1) -> str:
    return f"{100.0 * x:.{nd}f}%"


def set_run(run, size=12, bold=False, color=INK):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def text_box(slide, x, y, w, h, text="", size=12, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for para in tf.paragraphs:
        if align is not None:
            para.alignment = align
        for run in para.runs:
            set_run(run, size, bold, color)
    return box


def bullets(slide, x, y, w, h, items, size=11, color=INK, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(gap)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(0.7)
    return shp


def accent_bar(slide, x, y, w, color=CYAN):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.04))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()


def add_header(slide, title, subtitle=None, phase=None, color=CYAN):
    if phase:
        pill(slide, 0.55, 0.28, 1.35, 0.34, phase, color)
        tx = 2.03
    else:
        tx = 0.55
    text_box(slide, tx, 0.32, 9.6, 0.5, title, 22, True, NAVY)
    if subtitle:
        text_box(slide, tx, 0.88, 10.6, 0.3, subtitle, 9.8, False, MUTED)
    accent_bar(slide, tx, 1.27, 1.15, color)


def add_footer(slide, idx):
    text_box(slide, 0.55, 7.12, 7.7, 0.18, "Neural-operator surrogate for VBTS | phase process deck", 7.5, False, MUTED)
    text_box(slide, 12.05, 7.12, 0.7, 0.18, f"{idx:02d}", 7.5, True, MUTED, PP_ALIGN.RIGHT)


def pill(slide, x, y, w, h, text, color=BLUE, size=9):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    tf = shp.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    set_run(p.runs[0], size, True, WHITE)
    return shp


def metric(slide, x, y, w, label, value, note="", color=BLUE):
    rect(slide, x, y, w, 0.88, WHITE)
    text_box(slide, x + 0.12, y + 0.08, w - 0.24, 0.18, label.upper(), 7.6, True, MUTED)
    text_box(slide, x + 0.12, y + 0.33, w - 0.24, 0.28, value, 17, True, color)
    if note:
        text_box(slide, x + 0.12, y + 0.64, w - 0.24, 0.18, note, 7.4, False, MUTED)


def picture(slide, rel, x, y, w, h=None, border=True):
    path = ROOT / rel
    if not path.exists():
        rect(slide, x, y, w, h or 2.0, PALE_RED, RED)
        text_box(slide, x + 0.2, y + 0.2, w - 0.4, 0.35, f"Missing: {rel}", 10, True, RED)
        return None
    if border:
        rect(slide, x - 0.03, y - 0.03, w + 0.06, (h or 2.0) + 0.06, WHITE, LINE)
    if h is None:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def card(slide, x, y, w, h, title, items, color=BLUE, fill=WHITE):
    rect(slide, x, y, w, h, fill)
    accent_bar(slide, x + 0.16, y + 0.18, 0.75, color)
    text_box(slide, x + 0.16, y + 0.32, w - 0.32, 0.28, title, 12, True, NAVY)
    bullets(slide, x + 0.16, y + 0.72, w - 0.32, h - 0.82, items, 9.4, INK, 3)


def module(slide, x, y, w, h, title, subtitle="", color=BLUE, fill=WHITE, title_size=10):
    rect(slide, x, y, w, h, fill, color)
    text_box(slide, x + 0.08, y + 0.12, w - 0.16, 0.23, title, title_size, True, NAVY, PP_ALIGN.CENTER)
    if subtitle:
        text_box(slide, x + 0.10, y + 0.40, w - 0.20, h - 0.48, subtitle, 7.3, False, MUTED, PP_ALIGN.CENTER)


def arrow_text(slide, x, y, w=0.35, text="->"):
    text_box(slide, x, y, w, 0.22, text, 14, True, MUTED, PP_ALIGN.CENTER)


def field_icon(slide, x, y, cell=0.13, color=CYAN, accent=ORANGE):
    for r in range(5):
        for c in range(5):
            shp = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                Inches(x + c * cell),
                Inches(y + r * cell),
                Inches(cell * 0.82),
                Inches(cell * 0.82),
            )
            shp.fill.solid()
            if (r - 2) ** 2 + (c - 2) ** 2 <= 2:
                shp.fill.fore_color.rgb = accent
            else:
                shp.fill.fore_color.rgb = RGBColor(224, 234, 242)
            shp.line.color.rgb = WHITE


def table(slide, x, y, w, h, headers, rows, col_widths=None, font=8.4):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(w * cw / total)
    for j, head in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = Inches(0.04)
        cell.margin_right = Inches(0.04)
        cell.text = str(head)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.runs[0], font, True, WHITE)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(242, 246, 250)
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.text = str(val)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            set_run(p.runs[0], font, False, INK)
    return shape


def section_slide(prs, blank, title, subtitle, phase, color, idx_fn):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color
    text_box(s, 0.75, 0.72, 1.8, 0.35, phase.upper(), 13, True, WHITE)
    text_box(s, 0.72, 1.35, 10.3, 0.85, title, 34, True, WHITE)
    text_box(s, 0.76, 2.35, 9.0, 0.42, subtitle, 14, False, WHITE)
    accent_bar(s, 0.78, 3.05, 1.45, WHITE)
    text_box(s, 0.78, 6.95, 3.4, 0.2, f"{idx_fn():02d}", 8, True, WHITE)
    return s


def normal_slide(prs, blank, title, subtitle=None, phase=None, color=CYAN):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    add_header(s, title, subtitle, phase, color)
    add_footer(s, len(prs.slides))
    return s


def build():
    benchmark = load_json("runs/phase3_fem/benchmark.json")
    baselines = load_json("runs/phase3_fem/vbts_baselines.json")
    hybrid = load_json("runs/phase3_fem/hybrid_multiseed.json")
    control = load_json("runs/phase4/policy_servo.json")
    sensor_build = load_json("runs/phase5/sensor_build.json")
    sensor_compat = load_json("runs/phase5/sensor_compat.json")
    sensor_inv = load_json("runs/phase5/sensor_inverse.json")
    sensor_multi = load_json("runs/phase5/sensor_inverse_multiframe.json")
    env = load_json("runs/phase6/env_demo.json")
    temporal = load_json("runs/phase7/temporal.json")
    temporal_cmp = load_json("runs/phase7/temporal_compare.json")
    temporal_bc = load_json("runs/phase7_bc_corrected/temporal.json")
    temporal_bc_cmp = load_json("runs/phase7_bc_corrected/temporal_compare.json")
    geom = load_json("runs/phase3_fem/geometry_ood_phaseE.json")

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1. Cover
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    text_box(s, 0.72, 0.52, 2.6, 0.28, "PROCESS DECK", 10, True, CYAN)
    text_box(s, 0.68, 0.88, 7.15, 1.28, "Quá trình thực hiện nghiên cứu\nNeural Operator VBTS theo phase", 26, True, NAVY)
    text_box(s, 0.72, 2.28, 7.5, 0.45, "Differentiable Neural-Operator Surrogate for Vision-Based Tactile Sensing and Control", 13, False, MUTED)
    metric(s, 0.76, 3.1, 2.25, "Final GT", "2520", "corrected-BC frames", CYAN)
    metric(s, 3.25, 3.1, 2.25, "Main model", "LR-FNO", "field-to-field", BLUE)
    metric(s, 5.74, 3.1, 2.25, "Accuracy", "0.0606", "rel-L2 test", GREEN)
    metric(s, 8.23, 3.1, 2.25, "Speedup", "29,182x", "vs single IPC solve", ORANGE)
    picture(s, "runs/phase5/gt_vs_fno_samples.png", 8.45, 0.55, 3.95, 2.2)
    picture(s, "runs/phase3_fem/fidelity_speed.png", 8.58, 4.28, 3.72, 2.15)
    add_footer(s, len(prs.slides))

    # 2. Roadmap
    s = normal_slide(prs, blank, "Roadmap theo phase", "Luồng nghiên cứu: chọn ground truth -> học operator -> downstream -> audit vật lý -> mở rộng hình học", "map", NAVY)
    phase_rows = [
        ["P0", "Khảo sát khả thi", "scale gel, mesh, eps, marker sampling", "chọn cấu hình chạy production"],
        ["P1", "Sinh ground truth", "IPC/UIPC corrected-BC, K<=5", "dataset 2520 frame"],
        ["P2", "LR-FNO", "operator field-to-field", "model chính"],
        ["P3", "Benchmark", "RQ1-RQ3, baselines, multiseed", "headline accuracy/speed"],
        ["P4", "Control", "autograd vs ES", "sample efficiency"],
        ["P5", "Sensor", "marker renderer + image inverse", "render o LR-FNO"],
        ["P6", "Env", "one-step tactile environment", "gradient end-to-end"],
        ["P7", "BC audit", "rigid-shift root cause", "reground corrected-BC"],
        ["P8", "Geometry-OOD", "zero/few-shot 6 shapes", "scope + mitigation"],
    ]
    table(s, 0.75, 1.75, 11.85, 4.95, ["Phase", "Mục tiêu", "Việc làm", "Đầu ra"], phase_rows, [0.8, 2.2, 4.2, 2.5], 8.6)

    # Phase 0
    section_slide(prs, blank, "Phase 0", "Khảo sát realistic geometry và ổn định số trước production", "P0", CYAN, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 0: Chọn cấu hình mô phỏng", "Từ gel lớn/solver cũ sang thin-gel realistic và marker sampling ổn định", "P0", CYAN)
    card(s, 0.75, 1.75, 3.7, 2.3, "Thiết kế vật lý", [
        "Gel representative: 20 x 20 x 3 mm.",
        "Marker grid 32 x 32 trên top surface.",
        "Structured tetra mesh; production chọn res24.",
    ], CYAN, WHITE)
    card(s, 4.8, 1.75, 3.7, 2.3, "Solver knobs", [
        "eps_velocity = 2.5e-5.",
        "d_hat = 1e-4; contact resistance = 1e9.",
        "velocity_tol tightened để giảm stochasticity.",
    ], BLUE, WHITE)
    card(s, 8.85, 1.75, 3.7, 2.3, "Sampling", [
        "Nearest-neighbor top markers gây nhiễu theo mesh.",
        "Chuyển sang bilinear marker sampling.",
        "Provenance lưu vào NPZ để trace downstream.",
    ], GREEN, WHITE)
    table(s, 1.0, 4.65, 11.25, 1.55, ["Gate", "Kết quả", "Quyết định"], [
        ["res20 -> res24", "~12% drift", "res20 quá coarse"],
        ["res24 -> res28", "tangential 0.95%, normal 1.18%", "res24 đạt plateau"],
        ["K=3 smoke", "tangential noise ~1-2%", "đủ launch sweep"],
    ], [2.0, 4.5, 3.7], 8.5)

    s = normal_slide(prs, blank, "Phase 0: Eps-axis và smoke", "Bằng chứng chọn eps_velocity nhỏ hơn cấu hình ban đầu", "P0", CYAN)
    table(s, 0.85, 1.85, 5.8, 2.95, ["Eps step", "Tangential change", "Ý nghĩa"], [
        ["0.001 -> 0.0005", "29.21%", "quá thô"],
        ["0.0005 -> 0.00025", "19.14%", "chưa ổn"],
        ["0.00025 -> 0.0001", "10.27%", "còn drift"],
        ["0.0001 -> 0.00005", "4.97%", "gần hội tụ"],
        ["0.00005 -> 0.000025", "1.54%", "chọn production"],
    ], [2.0, 2.0, 2.6], 8.0)
    card(s, 7.0, 1.85, 5.2, 2.95, "Smoke production-shaped", [
        "Một frame random realistic: depth 0.614 mm, partial-slip.",
        "K=3 replicate noise: tangential 1.02%, normal 0.10%.",
        "Mean single solve ~11.8 s; K=3 target ~35.4 s.",
        "Kết luận: đủ ổn để chạy production, không overwrite GT cũ.",
    ], GREEN, PALE_GREEN)
    card(s, 0.85, 5.25, 11.35, 0.9, "Artifact chính", [
        "data/uipc/conv_realistic_eps_axis_vtol001/convergence_report.json; data/uipc/realistic_phase0/chosen_eps_regimes; infra/gen_uipc_sweep.sh",
    ], CYAN, WHITE)

    # Phase 1
    section_slide(prs, blank, "Phase 1", "Sinh IPC/UIPC ground truth realistic và corrected-BC", "P1", BLUE, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 1: Production ground truth", "Final dataset dùng cho paper và toàn bộ downstream", "P1", BLUE)
    metric(s, 0.78, 1.8, 2.25, "Frames", "2520", "2120 train / 400 test", BLUE)
    metric(s, 3.25, 1.8, 2.25, "Raw solves", "12,600", "63 x 40 x K=5", CYAN)
    metric(s, 5.72, 1.8, 2.25, "Mesh", "res24", "3125 verts, 13824 tets", GREEN)
    metric(s, 8.19, 1.8, 2.25, "Marker grid", "32 x 32", "bilinear sampling", ORANGE)
    metric(s, 10.66, 1.8, 1.65, "BC", "fixed", "bottom", RED)
    table(s, 0.95, 3.15, 5.55, 2.55, ["Mode", "All frames", "Test frames"], [
        ["normal", "425", "68"],
        ["stick", "712", "113"],
        ["partial-slip", "895", "142"],
        ["full-slip", "488", "77"],
    ], [2.2, 1.7, 1.7], 8.6)
    card(s, 7.0, 3.18, 4.9, 2.45, "Dataset source of truth", [
        "data/uipc/shear_res24_avg_swept_REALISTIC_BC.npz",
        "Provenance: gel_res, eps_velocity, velocity_tol, marker_sampling, solve_time_s.",
        "Downstream JSON đều trỏ về gt_path này.",
    ], BLUE, WHITE)

    s = normal_slide(prs, blank, "Phase 1: Robust averaging", "UIPC không bitwise-repeatable, nên target được lọc theo replicate distance", "P1", BLUE)
    table(s, 0.85, 1.75, 6.2, 2.85, ["Kept K", "Frames", "Ghi chú"], [
        ["5", "1397", "ổn định nhất"],
        ["4", "318", "lọc 1 outlier"],
        ["3", "331", "noise trung bình"],
        ["2", "474", "thường ở full-slip khó"],
    ], [1.5, 1.6, 4.0], 8.6)
    metric(s, 7.55, 1.85, 2.35, "Full-slip uniform", "22.0%", "mean energy", GREEN)
    metric(s, 10.15, 1.85, 2.35, "p95 uniform", "27.0%", "old soft-BC ~99.8%", ORANGE)
    table(s, 7.55, 3.25, 4.95, 1.65, ["Mode", "Tangential rep-noise"], [
        ["normal", "1.4%"],
        ["stick", "2.0%"],
        ["partial", "3.5%"],
        ["full", "4.4%"],
    ], [2.0, 2.7], 8.6)
    card(s, 0.85, 5.25, 11.65, 0.9, "Ý nghĩa", [
        "Phase 1 biến solver stochastic thành target đủ sạch cho operator learning, đồng thời giữ disclosure noise trong Limitations.",
    ], GREEN, PALE_GREEN)

    # Phase 2
    section_slide(prs, blank, "Phase 2", "Xây LR-FNO field-to-field surrogate", "P2", GREEN, lambda: len(prs.slides) + 1)

    s = normal_slide(prs, blank, "Phase 2: Neural Operator là gì?", "Học ánh xạ giữa các trường/hàm, không chỉ hồi quy từng điểm riêng lẻ", "P2", GREEN)
    rect(s, 0.9, 1.75, 11.45, 1.65, WHITE)
    field_icon(s, 1.25, 2.18, 0.13, CYAN, ORANGE)
    text_box(s, 1.15, 2.95, 1.2, 0.2, "input field a(x)", 7.5, True, MUTED, PP_ALIGN.CENTER)
    arrow_text(s, 2.55, 2.42, 0.55)
    module(s, 3.2, 2.05, 2.0, 0.82, "Neural Operator", "Gθ: a → u", GREEN, PALE_GREEN, 10)
    arrow_text(s, 5.42, 2.42, 0.55)
    field_icon(s, 6.15, 2.18, 0.13, GREEN, BLUE)
    text_box(s, 6.03, 2.95, 1.25, 0.2, "output field u(x)", 7.5, True, MUTED, PP_ALIGN.CENTER)
    module(s, 8.0, 2.05, 1.55, 0.82, "Train offline", "IPC/UIPC pairs", BLUE, PALE_BLUE, 9.4)
    arrow_text(s, 9.72, 2.42, 0.45)
    module(s, 10.25, 2.05, 1.55, 0.82, "Infer online", "ms latency", ORANGE, PALE_ORANGE, 9.4)
    card(s, 0.95, 4.05, 3.55, 1.55, "Khác MLP local", [
        "MLP dự đoán từng điểm dễ bỏ mất ngữ cảnh toàn cục.",
        "Neural operator học toàn bộ trường → toàn bộ trường.",
    ], CYAN, WHITE)
    card(s, 4.85, 4.05, 3.55, 1.55, "Vì sao hợp với VBTS", [
        "Gel bonded có đáp ứng đàn hồi non-local.",
        "Một contact cục bộ làm dịch chuyển cả lân cận.",
    ], GREEN, WHITE)
    card(s, 8.75, 4.05, 3.55, 1.55, "Vai trò trong nghiên cứu", [
        "Amortize solver cost: FEM/IPC sinh data offline.",
        "Surrogate frozen dùng cho control/sensor/env.",
    ], ORANGE, WHITE)

    s = normal_slide(prs, blank, "Phase 2: Kiến trúc FNO trunk", "Fourier Neural Operator dùng spectral convolution để học kernel toàn cục", "P2", GREEN)
    y = 2.05
    xs = [0.8, 2.35, 3.95, 5.65, 7.35, 9.05, 10.85]
    labels = [
        ("Input", "a(x): contact\nchannels"),
        ("Lift", "1x1 conv / MLP\nto latent width"),
        ("FFT", "spatial field →\nFourier modes"),
        ("Low modes", "keep k modes\nlearn R_theta"),
        ("IFFT + W", "global spectral\n+ pointwise path"),
        ("Activation", "GELU / stack\nL layers"),
        ("Project", "u_x,u_y,u_z\nmarker field"),
    ]
    colors = [CYAN, BLUE, PURPLE, ORANGE, GREEN, BLUE, CYAN]
    for i, ((title, sub), color) in enumerate(zip(labels, colors)):
        module(s, xs[i], y, 1.25 if i not in {3, 4} else 1.35, 0.95, title, sub, color, WHITE, 8.8)
        if i < len(xs) - 1:
            arrow_text(s, xs[i] + (1.32 if i not in {3, 4} else 1.43), y + 0.34, 0.28)
    rect(s, 3.72, 3.42, 5.05, 1.15, PALE_BLUE, BLUE)
    text_box(s, 3.92, 3.58, 4.65, 0.24, "FNO layer: v'(x) = σ( F⁻¹( Rθ · F(v) ) + Wv )", 11, True, NAVY, PP_ALIGN.CENTER)
    text_box(s, 4.08, 3.92, 4.35, 0.25, "Rθ học global kernel trong miền tần số; W giữ biến đổi local/pointwise.", 8.5, False, MUTED, PP_ALIGN.CENTER)
    card(s, 0.9, 5.0, 3.45, 1.38, "Điểm mạnh", [
        "Receptive field toàn cục ngay trong mỗi spectral layer.",
        "Rất hợp với response non-local của gel.",
    ], GREEN, PALE_GREEN)
    card(s, 4.95, 5.0, 3.45, 1.38, "Điểm yếu", [
        "Low modes làm trơn chi tiết near-contact.",
        "Cần refinement để bắt biên/sắc thái cục bộ.",
    ], ORANGE, PALE_ORANGE)
    card(s, 9.0, 5.0, 2.95, 1.38, "Trong deck", [
        "FNO trunk là ablation và nền cho LR-FNO.",
    ], BLUE, WHITE)

    s = normal_slide(prs, blank, "Phase 2: Kiến trúc LR-FNO đề xuất", "Locally Refined FNO = global spectral trunk + local refinement path", "P2", GREEN)
    module(s, 0.9, 2.05, 1.45, 0.78, "Input field", "depth + shear\n+ params", CYAN, WHITE)
    arrow_text(s, 2.48, 2.30, 0.32)
    module(s, 3.0, 1.55, 2.15, 0.8, "Global path", "FNO spectral trunk\ncaptures non-local response", BLUE, PALE_BLUE)
    module(s, 3.0, 3.0, 2.15, 0.8, "Local path", "2-layer conv stem\nkeeps near-contact features", GREEN, PALE_GREEN)
    arrow_text(s, 5.35, 1.88, 0.35)
    arrow_text(s, 5.35, 3.30, 0.35)
    module(s, 5.95, 2.25, 1.75, 0.9, "Fuse", "concat trunk +\nlocal + raw input", PURPLE, WHITE)
    arrow_text(s, 7.88, 2.58, 0.35)
    module(s, 8.45, 2.25, 1.95, 0.9, "Refinement head", "small conv head\ncorrects local residual", ORANGE, PALE_ORANGE)
    arrow_text(s, 10.58, 2.58, 0.35)
    module(s, 11.1, 2.25, 1.35, 0.9, "Output", "u_x,u_y,u_z", CYAN, WHITE)
    table(s, 1.0, 4.55, 5.35, 1.55, ["Component", "Vai trò"], [
        ["FNO trunk", "global / non-local elastic response"],
        ["Local stem", "near-contact texture and sharp gradients"],
        ["Refinement head", "residual correction with small parameter cost"],
    ], [1.8, 3.4], 8.2)
    table(s, 6.95, 4.55, 4.95, 1.55, ["Evidence", "Kết quả"], [
        ["5-seed rel-L2", "0.0590 +/- 0.0014"],
        ["vs FNO trunk", "wins 5/5 seeds, t=4.72"],
        ["Cost", "+3.3% params, lower fps"],
    ], [2.1, 2.6], 8.2)

    s = normal_slide(prs, blank, "Phase 2: Field-to-field framing", "Bài toán được đặt là học toán tử contact field -> marker displacement field", "P2", GREEN)
    rect(s, 0.85, 2.0, 11.6, 1.25, WHITE)
    for i, (txt, color) in enumerate([("Contact field", CYAN), ("FNO spectral trunk", BLUE), ("Local refinement", GREEN), ("Marker displacement", ORANGE)]):
        x = 1.05 + 2.8 * i
        pill(s, x, 2.45, 1.75, 0.38, txt, color, 8.5)
        if i < 3:
            text_box(s, x + 1.9, 2.46, 0.55, 0.25, "->", 15, True, MUTED, PP_ALIGN.CENTER)
    card(s, 0.9, 4.0, 3.55, 1.65, "Input", [
        "Indentation profile.",
        "Tangential-drive channel masked to contact.",
        "Physical params R, mu, E.",
    ], CYAN, WHITE)
    card(s, 4.85, 4.0, 3.55, 1.65, "Why FNO", [
        "Gel response is non-local.",
        "Spectral kernel sees global context.",
        "Per-point local MLP collapses on slip.",
    ], BLUE, WHITE)
    card(s, 8.8, 4.0, 3.55, 1.65, "Why LR", [
        "Local stem captures near-contact sharpness.",
        "Small refinement head.",
        "Chosen by multiseed stability.",
    ], GREEN, WHITE)

    s = normal_slide(prs, blank, "Phase 2: Model selection", "LR-FNO được chọn bằng five-seed benchmark, không chỉ seed-0 peak", "P2", GREEN)
    rows = []
    for key, name in [("unet", "U-Net"), ("fno", "FNO trunk"), ("ufno", "U-FNO-style"), ("local_refined_fno", "LR-FNO")]:
        m = hybrid["models"][key]
        rows.append([
            name,
            f"{m['rel_l2_overall']['mean']:.4f} +/- {m['rel_l2_overall']['std']:.4f}",
            f"{m['dir_deg']['mean']:.2f}",
            f"{m['throughput_fps']:.0f}",
            "---" if key == "fno" else f"{m.get('paired_vs_fno', {}).get('paired_t', 0):+.2f}",
        ])
    table(s, 0.85, 1.75, 11.6, 2.55, ["Model", "rel-L2 mean/std", "dir deg", "fps", "paired t vs FNO"], rows, [2.0, 2.6, 1.4, 1.6, 2.0], 8.2)
    card(s, 1.0, 4.75, 5.2, 1.45, "Quyết định", [
        "LR-FNO thắng FNO 5/5 seed, paired t=4.72, std nhỏ.",
        "U-FNO-style có seed tốt nhưng variance lớn; U-Net nhanh nhưng kém chính xác hơn.",
    ], GREEN, PALE_GREEN)
    card(s, 6.85, 4.75, 4.9, 1.45, "Framing khoa học", [
        "Gọi là locally refined FNO / U-FNO family adaptation.",
        "Không claim kiến trúc hoàn toàn mới.",
    ], BLUE, PALE_BLUE)

    # Phase 3
    section_slide(prs, blank, "Phase 3", "Benchmark RQ1-RQ3: fidelity, generalization, speed", "P3", ORANGE, lambda: len(prs.slides) + 1)
    rq1 = benchmark["RQ1"]["fno"]
    mlp = benchmark["RQ1"]["mlp"]
    rq3 = benchmark["RQ3"]
    s = normal_slide(prs, blank, "Phase 3: Headline metrics", "Số chính từ runs/phase3_fem/benchmark.json", "P3", ORANGE)
    metric(s, 0.8, 1.75, 2.25, "rel-L2", f"{rq1['relative_l2']['overall']:.4f}", "LR-FNO overall", BLUE)
    metric(s, 3.25, 1.75, 2.25, "Direction", f"{rq1['tangential_dir_error_deg']:.2f} deg", "tangential", CYAN)
    metric(s, 5.7, 1.75, 2.25, "MLP gap", f"{mlp['relative_l2']['overall'] / rq1['relative_l2']['overall']:.1f}x", "field-to-field wins", GREEN)
    metric(s, 8.15, 1.75, 2.25, "FPS", f"{rq3['throughput_fps']['fno']:.0f}", "operator inference", ORANGE)
    metric(s, 10.6, 1.75, 1.85, "Speedup", f"{rq3['fno_speedup_vs_gt_solver']:,.0f}x", "single solve", RED)
    table(s, 0.85, 3.05, 6.0, 2.2, ["Mode", "LR-FNO rel-L2", "MLP rel-L2"], [
        ["normal", f"{rq1['relative_l2']['normal']:.3f}", f"{mlp['relative_l2']['normal']:.3f}"],
        ["stick", f"{rq1['relative_l2']['stick']:.3f}", f"{mlp['relative_l2']['stick']:.3f}"],
        ["partial", f"{rq1['relative_l2']['partial_slip']:.3f}", f"{mlp['relative_l2']['partial_slip']:.3f}"],
        ["full", f"{rq1['relative_l2']['full_slip']:.3f}", f"{mlp['relative_l2']['full_slip']:.3f}"],
    ], [2.2, 2.0, 2.0], 8.5)
    picture(s, "runs/phase3_fem/fidelity_speed.png", 7.35, 3.1, 4.8, 2.55)

    s = normal_slide(prs, blank, "Phase 3: Baseline bake-off", "Bảng native PowerPoint, số từ vbts_baselines.json", "P3", ORANGE)
    base_rows = []
    for key, label in [
        ("tacto_kinematic", "TACTO-style"),
        ("cattaneo_mindlin_analytic", "Cattaneo-Mindlin"),
        ("taxim_fots_linear", "Taxim/FOTS linear"),
        ("mlp_perpoint", "MLP per-point"),
        ("deeponet", "DeepONet"),
        ("unet", "U-Net"),
        ("galerkin_transformer", "Galerkin"),
        ("fno_ours", "FNO trunk"),
        ("lr_fno_ours", "LR-FNO"),
    ]:
        m = baselines["models"][key]
        adv = baselines["fno_advantage_x"].get(key, "---")
        if key in {"fno_ours", "lr_fno_ours"}:
            adv = "---"
        adv_s = "---" if adv == "---" else f"{adv:.2f}x"
        base_rows.append([label, f"{m['relative_l2']['overall']:.3f}", f"{m['tangential_dir_error_deg']:.1f}", adv_s])
    table(s, 0.75, 1.65, 6.3, 4.95, ["Method", "rel-L2", "dir deg", "FNO adv."], base_rows, [3.2, 1.3, 1.3, 1.3], 7.8)
    card(s, 7.45, 1.8, 4.65, 2.1, "Kết luận", [
        "Classical/analytic/linear marker cores kém rõ trên cùng GT.",
        "Dense neural models tiến gần hơn; U-Net là đối thủ sát.",
        "Thông điệp đúng: field-to-field non-local operator learning là chìa khóa.",
    ], ORANGE, PALE_ORANGE)
    card(s, 7.45, 4.25, 4.65, 1.55, "Sắc thái", [
        "Không claim raw fps thắng mọi simulator.",
        "Lợi thế bền là latency thấp + autograd + mesh-free deployment.",
    ], BLUE, PALE_BLUE)

    s = normal_slide(prs, blank, "Phase 3: RQ2 và slip heads", "Generalization theo tham số vật lý và phân loại contact mode", "P3", ORANGE)
    rq2 = benchmark["RQ2"]
    table(s, 0.9, 1.8, 5.5, 2.0, ["OOD axis", "Degradation", "N OOD"], [
        ["high radius R", f"{rq2['high_radius']['degradation_x']:.2f}x", rq2["high_radius"]["n_ood"]],
        ["high friction mu", f"{rq2['high_mu']['degradation_x']:.2f}x", rq2["high_mu"]["n_ood"]],
        ["high stiffness E", f"{rq2['high_E']['degradation_x']:.2f}x", rq2["high_E"]["n_ood"]],
    ], [2.7, 1.5, 1.1], 8.6)
    sh = benchmark["RQ1"]["slip_head_a_multitask"]
    table(s, 7.05, 1.8, 5.05, 2.3, ["Slip head A", "F1"], [
        ["macro", f"{sh['macro_f1']:.3f}"],
        ["binary slip", f"{sh['slip_f1']:.3f}"],
        ["normal", f"{sh['per_class_f1']['normal']:.3f}"],
        ["full slip", f"{sh['per_class_f1']['full_slip']:.3f}"],
    ], [2.4, 1.4], 8.6)
    card(s, 1.0, 4.65, 10.85, 1.25, "Interpretation", [
        "Radius là trục nhạy nhất; E bền hơn trong displacement-control. Slip-head tốt nhưng vẫn báo rõ per-class thay vì chỉ macro.",
    ], GREEN, PALE_GREEN)

    # Phase 4
    section_slide(prs, blank, "Phase 4", "Học policy điều khiển qua surrogate khả vi", "P4", BLUE, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 4: Task design", "Servo policy một bước, không phải RL đa bước", "P4", BLUE)
    card(s, 0.85, 1.75, 3.7, 2.3, "Input", [
        "Context: mu, E, R, geometry.",
        "Target summary từ tactile field.",
        "PolicyMLP xuất action sx, sy.",
    ], CYAN, WHITE)
    card(s, 4.85, 1.75, 3.7, 2.3, "Autograd branch", [
        "LR-FNO frozen: eval + requires_grad False.",
        "Gradient chỉ đi tới action/policy.",
        "1 backward per step.",
    ], GREEN, WHITE)
    card(s, 8.85, 1.75, 3.7, 2.3, "ES baseline", [
        "OpenAI-ES antithetic.",
        "Same PolicyMLP and optimizer.",
        "Cô lập nguồn gradient.",
    ], ORANGE, WHITE)
    picture(s, "runs/phase4/policy_servo_curve.png", 1.0, 4.65, 5.35, 1.85)
    card(s, 6.75, 4.65, 5.2, 1.85, "Scope", [
        "Trong simulation, static one-step map.",
        "Dùng để test differentiability/sample-efficiency, không claim closed-loop dynamics.",
    ], BLUE, PALE_BLUE)

    s = normal_slide(prs, blank, "Phase 4: Autograd vs ES", "Số từ runs/phase4/policy_servo.json", "P4", BLUE)
    auto, es = control["autograd"], control["es"]
    table(s, 0.85, 1.75, 7.0, 2.35, ["Method", "Final loss", "Fwd queries", "Wall time", "Target query"], [
        ["Autograd", f"{auto['final_loss']['mean']:.2e}", auto["forward_evals"], f"{auto['wall_s']['mean']:.1f}s", "81"],
        ["ES pop32", f"{es['final_loss']['mean']:.2e}", es["forward_evals"], f"{es['wall_s']['mean']:.1f}s", "5824"],
    ], [1.8, 1.8, 1.5, 1.5, 1.5], 8.5)
    metric(s, 8.35, 1.82, 2.1, "Query", "72x", "fewer to target", GREEN)
    metric(s, 10.65, 1.82, 1.85, "Wall", "24x", "faster", CYAN)
    picture(s, "runs/phase4/policy_servo_curve.png", 1.0, 4.45, 5.4, 1.9)
    card(s, 7.05, 4.45, 4.95, 1.9, "Finding", [
        "Suh-style gradient pathology không xuất hiện ở one-step learned surrogate.",
        "Smooth operator gradient có ích cho control, dù không thay thế multi-step contact proof.",
    ], GREEN, PALE_GREEN)

    # Phase 5
    section_slide(prs, blank, "Phase 5", "Cảm biến marker-dot khả vi và image-space inverse", "P5", CYAN, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 5: Renderer pipeline", "Từ displacement field sang ảnh marker-dot bằng torch", "P5", CYAN)
    picture(s, "runs/phase5/preview.png", 0.85, 1.72, 5.8, 1.95)
    card(s, 7.05, 1.72, 4.95, 1.95, "Renderer", [
        "Pinhole camera nhìn từ dưới gel.",
        "Deformed markers -> Gaussian dot splatting.",
        "Torch implementation: autograd-connected.",
    ], CYAN, WHITE)
    rect(s, 0.9, 4.5, 11.3, 0.8, WHITE)
    for i, txt in enumerate(["FEM/IPC field", "marker positions", "rendered dots", "image loss", "gradient"]):
        pill(s, 1.05 + i * 2.15, 4.72, 1.35, 0.34, txt, [BLUE, CYAN, GREEN, ORANGE, RED][i], 7.4)
        if i < 4:
            text_box(s, 2.45 + i * 2.15, 4.72, 0.35, 0.22, "->", 12, True, MUTED)

    s = normal_slide(prs, blank, "Phase 5: Sensor metrics", "Số từ sensor_build / sensor_compat / sensor_inverse", "P5", CYAN)
    rt = sensor_build["round_trip"]
    table(s, 0.85, 1.72, 5.8, 2.45, ["Metric", "Value", "Note"], [
        ["flow-disp cosine", f"{sensor_build['flow_disp_cos_mean']:.3f}", "mean over modes"],
        ["round-trip overall", f"{rt['overall_px']:.2f} px", "render -> track"],
        ["stick / slip px", f"{rt['stick_px']:.2f} / {rt['slip_px']:.2f}", "mode split"],
        ["compat flow rel-L2", f"{sensor_compat['flow_rel_l2_overall']:.3f}", "render o LR-FNO"],
        ["single inverse", f"{100*sensor_inv['rel_err']:.1f}% / {sensor_inv['dir_err_deg']:.2f} deg", "full slip"],
    ], [2.5, 1.8, 2.5], 8.0)
    inv = sensor_multi["overall"]
    table(s, 7.1, 1.72, 4.95, 2.45, ["Mode", "Mag err", "Dir err"], [
        ["normal", "23.3%", "16.9 deg"],
        ["stick", "5.0%", "1.1 deg"],
        ["partial", "3.8%", "0.9 deg"],
        ["full", "15.3%", "0.4 deg"],
        ["overall", f"{inv['magnitude_error_pct']['mean']:.1f}%", f"{inv['direction_error_deg']['mean']:.1f} deg"],
    ], [1.7, 1.5, 1.5], 8.0)
    card(s, 0.95, 4.75, 11.0, 1.2, "Interpretation", [
        "Image inverse tốt ở shear-rich regimes; normal khó vì shear gần 0 khiến hướng ill-defined. Report phải tách theo mode.",
    ], GREEN, PALE_GREEN)

    s = normal_slide(prs, blank, "Phase 5: Visual evidence", "Ảnh lấy trực tiếp từ runs/phase5", "P5", CYAN)
    picture(s, "runs/phase5/gt_vs_fno_samples.png", 0.85, 1.65, 5.25, 4.45)
    picture(s, "runs/phase5/test_samples.png", 6.45, 1.65, 4.05, 4.45)
    card(s, 10.9, 1.85, 1.45, 3.35, "Read", [
        "GT/FNO residual.",
        "Flow residual.",
        "Mode samples.",
    ], CYAN, WHITE)

    # Phase 6
    section_slide(prs, blank, "Phase 6", "Framework mô phỏng VBTS và differentiable environment", "P6", GREEN, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 6: Env wrapper", "Frozen LR-FNO + marker sensor + image reward", "P6", GREEN)
    reward, grad = env["reward"], env["gradcheck"]
    metric(s, 0.85, 1.75, 2.35, "Reward gap", pct(reward["gap_closed_frac"]), "random -> oracle", GREEN)
    metric(s, 3.45, 1.75, 2.35, "Gradcheck", pct(grad["rel_error"]), "finite diff", CYAN)
    metric(s, 6.05, 1.75, 2.35, "Flows to action", "PASS", "autograd path", BLUE)
    metric(s, 8.65, 1.75, 2.35, "Obs image", "64 px", "noisy reward", ORANGE)
    picture(s, "runs/phase6/env_demo.png", 0.9, 3.25, 4.9, 2.6)
    card(s, 6.25, 3.25, 5.55, 2.6, "Scope", [
        "Contextual one-step environment, not multi-step dynamics.",
        "Useful proof: policy can optimize through image reward.",
        "Current formal evidence is gradient-flow sanity check.",
    ], GREEN, PALE_GREEN)

    s = normal_slide(prs, blank, "Phase 6: 6b-6d framework notes", "Các hướng framework còn lại hiện là log/report evidence, không còn đủ artifact trong runs/phase6", "P6", GREEN)
    table(s, 0.85, 1.75, 11.35, 3.3, ["Subphase", "Finding", "Evidence / caveat"], [
        ["6b realism", "camera noise floor không che được FNO error", "số từ log gốc; cần archive artifact nếu muốn reproducible"],
        ["6c temporal", "endpoint-only đủ; loading-history NULL", "model-free kNN matched: cross/same ~1.00"],
        ["6d object geometry", "sphere << cylinder < flat về marker-flow", "đặt nền cho Phase 8 geometry-OOD"],
        ["6a env", "artifact còn đầy đủ", "runs/phase6/env_demo.{json,png}"],
    ], [1.4, 4.2, 5.2], 8.2)
    card(s, 1.0, 5.45, 10.8, 0.8, "Audit note", [
        "Báo cáo Phase 6 đã ghi rõ chỉ 6a còn artifact trong runs; khi trình bày chính thức nên phân biệt measured artifact vs archived/log-derived evidence.",
    ], ORANGE, PALE_ORANGE)

    # Phase 7
    section_slide(prs, blank, "Phase 7", "Temporal visualization và audit boundary condition", "P7", RED, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 7: Soft-BC symptom", "FNO khớp GT soft-BC, nhưng GT có rigid drift artifact", "P7", RED)
    table(s, 0.85, 1.7, 5.7, 2.2, ["Temporal compare", "Value"], [
        ["mean rel-L2", f"{temporal_cmp['mean_rel_l2']:.3f}"],
        ["EPE", f"{temporal_cmp['sensor_summary']['mean_epe_px']:.2f} px"],
        ["dir error", f"{temporal_cmp['sensor_summary']['mean_dir_deg']:.2f} deg"],
        ["cos", f"{temporal_cmp['sensor_summary']['mean_cos']:.4f}"],
    ], [3.0, 2.0], 8.6)
    picture(s, "runs/phase7/temporal_gt_vs_fno.png", 7.0, 1.65, 4.8, 3.95)
    card(s, 0.95, 4.55, 5.35, 1.25, "Interpretation", [
        "Soft-BC run chứng minh model khớp lineage cũ, nhưng không phải evidence vật lý cuối vì bottom constraint tạo rigid shift.",
    ], RED, PALE_RED)

    s = normal_slide(prs, blank, "Phase 7: Corrected-BC root cause", "Fixed bottom loại bỏ near-rigid tangential drift", "P7", RED)
    table(s, 0.85, 1.7, 5.9, 2.3, ["Field diagnostic", "Soft-BC", "Corrected-BC"], [
        ["local normal contrast", "0.085 mm", "0.258-0.262 mm"],
        ["residual tangential", "0.012 mm", "0.076-0.090 mm"],
        ["mean marker flow", "3.1-3.3 px", "0.54-0.67 px"],
        ["uniform energy", "~99.8%", "20-24% smoke"],
    ], [2.5, 1.6, 1.8], 8.0)
    picture(s, "runs/phase7_bc_corrected/phase7_field_diagnostic.png", 7.1, 1.75, 4.85, 1.55)
    picture(s, "runs/phase7_bc_corrected/temporal_video.png", 1.0, 4.55, 5.5, 1.75)
    picture(s, "runs/phase7/temporal_video.png", 7.0, 4.55, 5.5, 1.75)

    s = normal_slide(prs, blank, "Phase 7: Evidence hygiene", "Không dùng corrected-BC GT-vs-FNO cũ làm claim model", "P7", RED)
    table(s, 0.9, 1.75, 6.2, 2.25, ["Artifact", "Status", "Use"], [
        ["runs/phase7/*", "soft-BC", "temporal visualization only"],
        ["runs/phase7_bc_corrected/temporal*", "corrected GT", "GT-only evidence valid"],
        ["runs/phase7_bc_corrected/temporal_gt_vs_fno", "FNO trained on old GT", "do not use as FNO evidence"],
        ["production REALISTIC_BC", "rerun complete", "source for Phase 3-6"],
    ], [2.5, 1.6, 3.0], 7.8)
    metric(s, 7.55, 1.82, 2.2, "Bad compare", "3.37", "rel-L2", RED)
    metric(s, 10.0, 1.82, 2.2, "Bad EPE", "2.77px", "old model vs new BC", ORANGE)
    card(s, 7.55, 3.25, 4.65, 1.75, "Outcome", [
        "Boundary-condition audit triggered full production reground.",
        "All final downstream metrics use corrected-BC dataset.",
    ], GREEN, PALE_GREEN)

    # Phase 8
    section_slide(prs, blank, "Phase 8", "Geometry-OOD: đo zero-shot drop và few-shot recovery", "P8", PURPLE, lambda: len(prs.slides) + 1)
    s = normal_slide(prs, blank, "Phase 8: Protocol", "LR-FNO train trên sphere, test trên 6 hình lệch dần khỏi sphere", "P8", PURPLE)
    lr = geom["models"]["lr_fno"]["aggregate"]
    id_rel = lr["in_distribution"]["relative_l2"]["overall"]
    id_um = lr["in_distribution"]["abs_rmse_um"]["overall"]["tangential"]
    metric(s, 0.9, 1.75, 2.4, "In-dist rel-L2", f"{id_rel['mean']:.4f}", f"+/- {id_rel['std']:.4f}", BLUE)
    metric(s, 3.55, 1.75, 2.4, "Tangential RMSE", f"{id_um['mean']:.2f}um", f"+/- {id_um['std']:.2f}", CYAN)
    metric(s, 6.2, 1.75, 2.4, "OOD shapes", "6", "4 analytic + 2 mesh", PURPLE)
    metric(s, 8.85, 1.75, 2.4, "Few-shot", "50", "frames, 20 epochs", GREEN)
    card(s, 0.95, 3.35, 5.4, 1.8, "Three conditions", [
        "Zero-shot: áp model sphere trực tiếp.",
        "Few-shot: finetune 50 OOD frames, lr=1e-4.",
        "Scratch: random init với cùng 50-frame budget.",
    ], PURPLE, WHITE)
    card(s, 6.85, 3.35, 5.0, 1.8, "Metric choice", [
        "Figure dùng tangential abs-RMSE um.",
        "rel-L2 vẫn lưu trong JSON nhưng phồng khi target norm nhỏ, đặc biệt ellipsoid.",
    ], ORANGE, PALE_ORANGE)

    s = normal_slide(prs, blank, "Phase 8: OOD result table", "Bảng native PowerPoint, số từ geometry_ood_phaseE.json", "P8", PURPLE)
    order = ["sphere_oodR", "rounded_tip", "bolt_hex", "cylinder", "cuboid", "ellipsoid"]
    names = {"sphere_oodR": "sphere_oodR", "rounded_tip": "rounded_tip", "bolt_hex": "bolt_hex", "cylinder": "cylinder", "cuboid": "cuboid", "ellipsoid": "ellipsoid"}
    rows = []
    for key in order:
        d = lr["datasets"][key]
        def t(stage):
            m = d[stage]["abs_rmse_um"]["overall"]["tangential"]
            return f"{m['mean']:.1f} +/- {m['std']:.1f}"
        z = d["zero_shot"]["abs_rmse_um"]["overall"]["tangential"]["mean"]
        f = d["few_shot"]["abs_rmse_um"]["overall"]["tangential"]["mean"]
        rec = (f / z - 1.0) * 100.0
        rows.append([names[key], f"{d['geometry_distance']['mean']:.2f}", t("zero_shot"), t("few_shot"), t("scratch"), f"{rec:.0f}%"])
    table(s, 0.65, 1.62, 12.05, 4.2, ["Shape", "dist", "zero um", "few um", "scratch um", "few delta"], rows, [2.1, 0.8, 2.0, 2.0, 2.0, 1.2], 7.7)

    s = normal_slide(prs, blank, "Phase 8: Figure and limitations", "Geometry-OOD là điểm yếu được đo và giảm nhẹ, không phải đã giải xong tổng quát hình học", "P8", PURPLE)
    picture(s, "docs/kse2026/figs/geometry_ood.png", 0.85, 1.65, 6.2, 3.95)
    card(s, 7.45, 1.7, 4.75, 1.5, "Findings", [
        "Zero-shot degrade tăng theo geometry distance.",
        "Few-shot thắng scratch cùng budget.",
        "Finetune tạo specialist, có forgetting trên sphere.",
    ], PURPLE, WHITE)
    card(s, 7.45, 3.55, 4.75, 1.45, "Scope", [
        "Không claim distant curved geometry.",
        "Rounded-tip là curved-bottom mesh PoC gần sphere.",
        "Next: multi-geometry foundation/adapters.",
    ], ORANGE, PALE_ORANGE)

    # Close
    s = normal_slide(prs, blank, "Tổng kết", "Các phase tạo thành một câu chuyện thống nhất: physics-grounded, differentiable, honest about scope", "wrap", NAVY)
    table(s, 0.85, 1.75, 11.65, 3.15, ["Layer", "Đóng góp", "Bằng chứng"], [
        ["Ground truth", "IPC/UIPC corrected-BC realistic thin gel", "2520 frames, robust adaptive-K"],
        ["Surrogate", "LR-FNO field-to-field", "rel-L2 0.0606, 29,182x vs IPC solve"],
        ["Downstream", "control + sensor + env", "72x fewer target queries, image inverse, reward gap 99.8%"],
        ["Audit", "BC root-cause and reground", "uniform energy old 99.8% -> corrected 20-24% smoke"],
        ["Generalization", "geometry-OOD quantified", "6 shapes, few-shot recovery 20-60%"],
    ], [2.0, 4.4, 5.0], 8.0)
    card(s, 1.05, 5.35, 10.95, 0.85, "Hướng tiếp theo", [
        "Multi-geometry foundation model, adapter/few-shot không quên, và sim-to-real calibration trên cảm biến marker-gel thật.",
    ], GREEN, PALE_GREEN)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    print(f"slides={len(prs.slides)}")


if __name__ == "__main__":
    build()
